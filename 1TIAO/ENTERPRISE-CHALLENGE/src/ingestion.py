"""Entrada limitada a tabelas CSV/XLSX; não interpreta PDF/OCR."""
from io import BytesIO, StringIO
from pathlib import Path
import csv
import zipfile
import pandas as pd
from openpyxl import load_workbook
from utils import normalize

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000
MAX_COLS = 50

def read_table(content: bytes, filename: str) -> pd.DataFrame:
    if not content or len(content) > MAX_BYTES:
        raise ValueError('Arquivo vazio ou acima de 5 MB.')
    ext = Path(filename).suffix.lower()
    try:
        if ext == '.csv':
            text = content.decode('utf-8-sig')
            first = text.splitlines()[0]
            delimiter = ';' if first.count(';') > first.count(',') else ','
            rows = list(csv.reader(StringIO(text), delimiter=delimiter))
            rows = [r for r in rows if any(str(v).strip() for v in r)]
        elif ext == '.xlsx':
            with zipfile.ZipFile(BytesIO(content)) as z:
                if sum(i.file_size for i in z.infolist()) > 40 * 1024 * 1024:
                    raise ValueError('XLSX excede 40 MB de conteúdo expandido.')
                if any(i.flag_bits & 1 for i in z.infolist()):
                    raise ValueError('Arquivo criptografado não suportado.')
            book = load_workbook(BytesIO(content), read_only=True, data_only=False)
            try:
                sheet = book.worksheets[0]
                if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                    raise ValueError('Máximo de 10.000 linhas de dados.')
                if sheet.max_column and sheet.max_column > MAX_COLS:
                    raise ValueError('Máximo de 50 colunas.')
                rows = []
                for row in sheet.iter_rows():
                    if any(c.data_type == 'f' for c in row):
                        raise ValueError('Use valores, não fórmulas, no XLSX de entrada.')
                    vals = [c.value for c in row]
                    if any(v is not None and str(v).strip() for v in vals):
                        rows.append(vals)
                    if len(rows) > MAX_ROWS + 1:
                        raise ValueError('Máximo de 10.000 linhas de dados.')
            finally:
                book.close()
        else:
            raise ValueError('Formato não suportado. Envie CSV ou XLSX.')
        if len(rows) < 2:
            raise ValueError('Arquivo precisa de cabeçalho e ao menos uma linha de dados.')
        header = [normalize(v) if v is not None else '' for v in rows[0]]
        if len(header) > MAX_COLS or len(rows)-1 > MAX_ROWS:
            raise ValueError('Limite: 10.000 linhas e 50 colunas.')
        if any(not h for h in header) or len(set(header)) != len(header):
            raise ValueError('Cabeçalhos vazios ou duplicados após normalização.')
        if any(len(row) != len(header) for row in rows[1:]):
            raise ValueError('Quantidade de campos inconsistente entre linhas.')
        return pd.DataFrame(rows[1:], columns=header)
    except (UnicodeError, zipfile.BadZipFile, csv.Error, KeyError, IndexError) as exc:
        raise ValueError('Arquivo ilegível ou formato inválido; confira o modelo.') from exc
